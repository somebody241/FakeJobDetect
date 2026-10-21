import hashlib
import importlib
import io
import re
import pandas as pd
import pandasql as ps
import pdfplumber
import streamlit as st
from docx import Document
import train

importlib.reload(train)

LABELS = {
    "salary_min": "Minimum salary",
    "salary_max": "Maximum salary",
    "required_experience_years": "Experience required (years)",
    "skills_listed": "Skills listed",
    "allcaps_ratio": "Capital-letter ratio",
    "posting_hour": "Hour posted",
    "company_age_months": "Company age (months)",
    "company_size": "Company size",
    "mentions_trial_or_payment": "Asks for a fee or payment details",
    "no_pay_range": "No pay range published",
}
ALIASES = {
    "salary_min": ["salary_min", "salary min", "minimum salary", "min salary"],
    "salary_max": ["salary_max", "salary max", "maximum salary", "max salary"],
    "required_experience_years": [
        "required_experience_years",
        "experience years",
        "years of experience",
        "experience required",
    ],
    "skills_listed": ["skills_listed", "skills listed", "number of skills"],
    "allcaps_ratio": ["allcaps_ratio", "capital ratio", "all caps ratio"],
    "posting_hour": ["posting_hour", "posting hour", "hour posted"],
    "company_age_months": ["company_age_months", "company age", "company age months"],
    "company_size": ["company_size", "company size"],
    "mentions_trial_or_payment": [
        "mentions_trial_or_payment",
        "trial or payment",
        "asks for payment",
    ],
    "no_pay_range": ["no pay range", "salary missing"],
}
INT_KEYS = {"skills_listed", "posting_hour", "company_size"}
FLOAT_KEYS = {
    "salary_min",
    "salary_max",
    "required_experience_years",
    "allcaps_ratio",
    "company_age_months",
}
BOOL_KEYS = {"no_pay_range", "mentions_trial_or_payment"}
MANUAL_FIELDS = ("company_age_months", "company_size", "posting_hour")

# Text parsing regexes for extracting the features from a file
PAYMENT_RE = re.compile(
    r"trial fee|deposit|payment details|upfront|wire transfer|gift card|"
    r"processing fee|training fee|bank details",
    re.I,
)
SALARY_RE = re.compile(
    r"(salary|pay|compensation)[^0-9]{0,40}\$?([\d,]+)k?\s*(-|to)\s*\$?([\d,]+)k?",
    re.I,
)
EXPERIENCE_RE = re.compile(
    r"(\d+(?:\.\d+)?)\s*\+?\s*years?(?: of)? experience",
    re.I,
)
NO_PAY_RE = re.compile(
    r"salary not (?:disclosed|listed|published)|no (?:salary|pay range)|unpaid",
    re.I,
)

ALIAS_TO_FIELD = {
    re.sub(r"[^a-z0-9]+", " ", alias.lower()).strip(): field
    for field, aliases in ALIASES.items()
    for alias in aliases
}

# Normalize text by removing special characters and converting to lowercase
def normalize(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()

# Example queries for the SQL interface
QUERIES = {
    "Scam rate by skills listed": (
        "SELECT skills_listed,\n"
        "       AVG(scam) AS scam_rate,\n"
        "       COUNT(*) AS postings\n"
        "FROM dataset\n"
        "GROUP BY skills_listed\n"
        "ORDER BY skills_listed"
    ),
    "Scam rate by exclamation marks": (
        "SELECT exclamation_marks,\n"
        "       AVG(scam) AS scam_rate,\n"
        "       COUNT(*) AS postings\n"
        "FROM dataset\n"
        "GROUP BY exclamation_marks\n"
        "ORDER BY exclamation_marks"
    ),
    "Scam rate by remote flag": (
        "SELECT is_remote,\n"
        "       AVG(scam) AS scam_rate,\n"
        "       COUNT(*) AS postings\n"
        "FROM dataset\n"
        "GROUP BY is_remote"
    ),
}

# my precious, cache🙏
# loads the model into memory and keeps it there
@st.cache_resource
def cached_saved_model(saved_at):
    return train.TheModel.load(train.MODEL_PATH)


# Set the starting values for the input form
def form_defaults():
    return {
        "salary_min": None,
        "salary_max": None,
        "no_pay_range": False,
        "required_experience_years": None,
        "skills_listed": None,
        "allcaps_ratio": None,
        "posting_hour": None,
        "company_age_months": None,
        "company_size": None,
        "mentions_trial_or_payment": False,
    }


# Read the text and tables from a PDF file
def read_pdf(data):
    texts = []
    tables = []
    with pdfplumber.open(io.BytesIO(data)) as document:
        for page in document.pages:
            texts.append(page.extract_text() or "")
            tables.extend(table for table in (page.extract_tables() or []) if table)
    return "\n".join(texts), tables


# Read the text and tables from a Word file
def read_docx(data):
    document = Document(io.BytesIO(data))
    texts = [paragraph.text for paragraph in document.paragraphs]
    tables = []
    for table in document.tables:
        tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
    return "\n".join(texts), tables


# Convert a raw value into the correct type for its feature
def parse_value(field, raw):
    text = str(raw).strip()
    if text.lower() in {"", "na", "n/a", "none", "nan", "-"}:
        return None
    if field in BOOL_KEYS:
        return text.lower() in {"1", "true", "yes", "y"}
    cleaned = text.replace(",", "").replace("$", "").strip()
    try:
        number = float(cleaned)
    except ValueError:
        return None
    if field in INT_KEYS:
        return int(number)
    return number



# Absorb a key-value pair into the found dictionary
def absorb_pair(found, key, raw):
    field = ALIAS_TO_FIELD.get(normalize(key))
    if field is None:
        return
    value = parse_value(field, raw)
    if value is not None:
        found[field] = value


# Extract structured features from the text and tables of a job posting
def structured_features(text, tables):
    found = {}
    for table in tables:
        if not table or not table[0]:
            continue
        headers = [normalize(cell or "") for cell in table[0]]
        fields = [ALIAS_TO_FIELD.get(header) for header in headers]
        if sum(field is not None for field in fields) >= 2 and len(table) > 1:
            for field, cell in zip(fields, table[1]):
                if field is not None:
                    value = parse_value(field, cell)
                    if value is not None:
                        found[field] = value
            break
        if len(table[0]) == 2:
            for row in table:
                if len(row) >= 2:
                    absorb_pair(found, row[0], row[1])
            if found:
                break
    for line in text.splitlines():
        separator = ":" if ":" in line else "=" if "=" in line else ""
        if not separator:
            continue
        key, _, raw = line.partition(separator)
        absorb_pair(found, key, raw)
    return found


# Find features mentioned in normal sentences and unstructured text
def prose_features(text):
    if not text.strip():
        return {}
    letters = [char for char in text if char.isalpha()]
    found = {
        "allcaps_ratio": (sum(char.isupper() for char in letters) / len(letters)) if letters else 0.0,
        "mentions_trial_or_payment": bool(PAYMENT_RE.search(text)),
    }
    if NO_PAY_RE.search(text):
        found["no_pay_range"] = True
    salary = SALARY_RE.search(text)
    if salary:
        minimum = float(salary.group(2).replace(",", ""))
        maximum = float(salary.group(4).replace(",", ""))
        if "k" in salary.group(0).lower() and maximum < 1000:
            minimum *= 1000
            maximum *= 1000
        if minimum >= 100 and maximum >= minimum:
            found["salary_min"] = minimum
            found["salary_max"] = maximum
            found["no_pay_range"] = False
    experience = EXPERIENCE_RE.search(text)
    if experience:
        found["required_experience_years"] = float(experience.group(1))
    skills = re.search(
        r"(?:skills|requirements|qualifications)\s*[:\-]\s*(.+)",
        text,
        re.I,
    )
    if skills:
        parts = [part.strip() for part in re.split(r",|\n|•", skills.group(1)) if part.strip()]
        if parts:
            found["skills_listed"] = len(parts)
    return found


# Extract features from an uploaded file
def extract_upload(name, data):
    lower = name.lower()
    if lower.endswith(".pdf"):
        text, tables = read_pdf(data)
    elif lower.endswith(".docx"):
        text, tables = read_docx(data)
    else:
        raise ValueError("Upload a PDF or a .docx file.")
    found = prose_features(text)
    found.update(structured_features(text, tables))
    if not text.strip() and not found:
        return {}, ["No text could be read from that file. Fill the fields below instead."]
    filled = [
        LABELS.get(key, key)
        for key, value in found.items()
        if not (key == "no_pay_range" and not value)
    ]
    notes = [f"Filled from the file: {', '.join(filled)}."]
    missing = [LABELS[key] for key in MANUAL_FIELDS if key not in found]
    if missing:
        notes.append("The file did not include: " + ", ".join(missing) + ".")
    return found, notes


# Clamp a value to the specified range
def clamp(key, value):
    limits = {
        "salary_min": (0, None),
        "salary_max": (0, None),
        "required_experience_years": (0, 80),
        "skills_listed": (0, 100),
        "allcaps_ratio": (0, 1),
        "posting_hour": (0, 23),
        "company_age_months": (0, 2000),
        "company_size": (1, 100000),
    }
    low, high = limits[key]
    value = max(low, value)
    if high is not None:
        value = min(high, value)
    return value


# Absorb the extracted features into the Streamlit session state and update the filled note
def apply_extracted(found, notes):
    for key, value in found.items():
        if key in BOOL_KEYS:
            st.session_state[key] = bool(value)
        elif key in INT_KEYS and value is not None:
            st.session_state[key] = int(clamp(key, value))
        elif key in FLOAT_KEYS and value is not None:
            st.session_state[key] = float(clamp(key, value))
    st.session_state.filled_note = " ".join(notes)
    if found:
        st.session_state.verdict = None
        st.session_state.probability = None
        st.session_state.closest_table = None
        st.session_state.pattern_body = ""
        st.session_state.result_text = (
            "Fields were filled from the file. Check this posting to see the verdict."
        )


# Check the form and create a DataFrame from the filled values
def posting_frame():
    no_pay = bool(st.session_state.no_pay_range)
    required = [
        "required_experience_years",
        "skills_listed",
        "allcaps_ratio",
        "posting_hour",
        "company_age_months",
        "company_size",
    ]
    if not no_pay:
        required = ["salary_min", "salary_max", *required]
    if any(st.session_state.get(key) is None for key in required):
        return None
    row = {
        "salary_min": None if no_pay else float(st.session_state.salary_min),
        "salary_max": None if no_pay else float(st.session_state.salary_max),
        "required_experience_years": float(st.session_state.required_experience_years),
        "skills_listed": int(st.session_state.skills_listed),
        "allcaps_ratio": float(st.session_state.allcaps_ratio),
        "posting_hour": int(st.session_state.posting_hour),
        "company_age_months": float(st.session_state.company_age_months),
        "company_size": int(st.session_state.company_size),
        "mentions_trial_or_payment": int(bool(st.session_state.mentions_trial_or_payment)),
    }
    return pd.DataFrame([row])


# Calculate how much each feature pushes the model toward a scam verdict
def feature_rows(model, posting):
    encoded, _ = train.prepare_features(posting, model.medians, model.columns)
    scaled = model._scaled(encoded).iloc[0]
    coefficients = model.clf.coef_[0]
    rows = []
    for name, coefficient, scaled_value in zip(model.columns, coefficients, scaled):
        push = float(scaled_value) * float(coefficient)
        if abs(push) < 1e-9:
            continue
        rows.append((push, name))
    return rows


# Format one feature value for displaying to the user
def display_value(name, posting):
    row = posting.iloc[0]
    if name == "mentions_trial_or_payment":
        return "yes" if int(row[name]) == 1 else "no"
    value = row[name]
    if pd.isna(value):
        return "not published"
    if name == "allcaps_ratio":
        return f"{float(value):.3f}"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.1f}"
    return str(value)


# Convert an internal feature name into its user-facing label
def feature_label(name):
    return LABELS.get(name, name)


# Create the lines that explain which features raise or lower the scam probability
def pattern_lines(model, posting):
    rows = feature_rows(model, posting)
    if pd.isna(posting.iloc[0]["salary_min"]):
        rows = [row for row in rows if row[1] not in {"salary_min", "salary_max"}]
    raises = [row for row in sorted(rows, reverse=True) if row[0] >= 0.3]
    lowers = [row for row in sorted(rows) if row[0] <= -0.3]
    lines = ["Feature pattern for this posting:"]
    lines.append("Raises the scam probability:" if raises else "No trait pushes this posting toward scam on its own.")
    lines.extend(f"- {feature_label(name)}: {display_value(name, posting)}" for _, name in raises)
    if lowers:
        lines.append("Lowers the scam probability:")
        lines.extend(f"- {feature_label(name)}: {display_value(name, posting)}" for _, name in lowers)
    return lines


# Create the text summary shown after a posting is checked
def result_text(verdict, probability, closest, lines):
    label = "scam" if verdict == 1 else "legitimate"
    written = [
        f"scam = {verdict}",
        f"Verdict: {label}",
        f"Scam probability: {probability:.1%}",
        "",
        "Three closest reviewed postings:",
    ]
    for number, (_, row) in enumerate(closest.iterrows(), start=1):
        written.append(
            f"{number}. scam = {int(row['scam'])}, skills {int(row['skills_listed'])}, "
            f"distance {row['distance']:.2f}, company age {row['company_age_months']:.0f} months, "
            f"capital ratio {row['allcaps_ratio']:.3f}"
        )
    written.append("")
    written.extend(lines)
    return "\n".join(written)


# the crown jewel, the scam and legitimate verdict panel
def verdict_panel(verdict, probability):
    # Builds one side of the verdict panel
    def side(title, detail, lit, color):
        background = color if lit else "#f4f4f5"
        foreground = "#ffffff" if lit else "#71717a"
        return (
            f'<div style="flex:1; background:{background}; color:{foreground}; '
            f'padding:28px 16px; text-align:center;">'
            f'<div style="font-size:14px; font-weight:700; letter-spacing:0.08em; '
            f'text-transform:uppercase;">{title}</div>'
            f'<div style="font-size:22px; font-weight:650; margin-top:8px;">{detail}</div>'
            f"</div>"
        )

    scam = verdict == 1
    legit = verdict == 0
    scam_detail = f"{probability:.1%}" if scam else "scam = 1"
    legit_detail = f"{probability:.1%}" if legit else "scam = 0"
    html = (
        '<div style="display:flex; border:1px solid #e4e4e7; border-radius:12px; '
        'overflow:hidden; margin:4px 0 16px 0;">'
        + side("Scam", scam_detail, scam, "#dc2626")
        + '<div style="width:1px; background:#e4e4e7;"></div>'
        + side("No scam", legit_detail, legit, "#16a34a")
        + "</div>"
    )
    st.markdown(html, unsafe_allow_html=True)


# Format the closest postings for display in a table
def neighbor_table(closest, features):
    show = closest[["distance", "scam", *features]].copy()
    show["distance"] = show["distance"].round(2)
    show["scam"] = show["scam"].astype(int)
    show["allcaps_ratio"] = show["allcaps_ratio"].round(3)
    return show


# Display the reviewed postings, statistics, charts, and filters
def render_backlog(reviews, scores, features):
    st.subheader("Reviewed backlog")
    st.write(
        "These 400 postings were already reviewed. scam = 1 is a confirmed scam and "
        "scam = 0 is legitimate. Held-out accuracy is the share of postings the model "
        "labels correctly when that posting was not in its training rows."
    )
    left, middle, right, far = st.columns(4)
    left.metric("Reviewed postings", f"{len(reviews)}")
    middle.metric("Confirmed scams", f"{int(reviews['scam'].sum())}")
    right.metric("Held-out accuracy", f"{scores['accuracy']:.1%}")
    far.metric("Scam recall", f"{scores['recall']:.1%}")
    st.caption(
        f"Held-out accuracy is {scores['accuracy']:.1%}. "
        f"Precision on scams is {scores['precision']:.1%}. The target for accuracy is 80%."
    )

    st.subheader("Characteristics associated with confirmed scams")
    st.write(
        "The model keeps a column only when its correlation with scam is at least 0.2. "
        "Category, vague words, exclamation marks, response delay, remote, and free email "
        "are left out. Bars above zero are higher in confirmed scams. Bars below zero are "
        "higher in legitimate postings."
    )
    numeric = list(features)
    correlation = reviews[numeric + ["scam"]].corr(numeric_only=True)["scam"].drop("scam")
    correlation = correlation.sort_values(ascending=False)
    correlation.index = [LABELS.get(name, name) for name in correlation.index]
    st.bar_chart(correlation)

    yes = reviews.loc[reviews["mentions_trial_or_payment"] == 1, "scam"].mean()
    no = reviews.loc[reviews["mentions_trial_or_payment"] == 0, "scam"].mean()
    st.dataframe(
        pd.DataFrame(
            [
                {
                    "Characteristic": LABELS["mentions_trial_or_payment"],
                    "Scam rate when yes": f"{yes:.0%}",
                    "Scam rate when no": f"{no:.0%}",
                }
            ]
        ),
        hide_index=True,
    )

    means = reviews.groupby("scam")[numeric].mean().T
    means.columns = ["Legitimate (scam = 0)", "Confirmed scam (scam = 1)"]
    means["gap"] = means.iloc[:, 1] - means.iloc[:, 0]
    means = means.sort_values("gap", ascending=False).drop(columns="gap")
    means.index = [LABELS.get(name, name) for name in means.index]
    st.write("Average value of each characteristic in the two reviewed groups.")
    st.dataframe(means.round(2))

    st.subheader("Filter the listings")
    categories = ["All", *sorted(reviews["category"].dropna().unique())]
    category = st.selectbox("Category", categories)
    youngest = int(reviews["company_age_months"].min())
    oldest = int(reviews["company_age_months"].max())
    age = st.slider("Company age (months)", youngest, oldest, (youngest, oldest))
    choice = st.radio(
        "Listings",
        ["All reviewed postings", "Confirmed scams", "Legitimate postings"],
        horizontal=True,
    )
    view = reviews[reviews["company_age_months"].between(age[0], age[1])]
    if category != "All":
        view = view[view["category"] == category]
    if choice == "Confirmed scams":
        view = view[view["scam"] == 1]
    elif choice == "Legitimate postings":
        view = view[view["scam"] == 0]
    st.write(f"{len(view)} postings match these filters.")
    if choice == "All reviewed postings" and not view.empty:
        rates = (view.groupby("category")["scam"].mean() * 100).sort_values(ascending=False)
        rates.name = "Confirmed scams (%)"
        st.write("Scam rate by category inside the current age filter.")
        st.bar_chart(rates)
    st.dataframe(view[["category", *features, "scam"]], hide_index=True)


# Turn a query result into chart data when possible
def chartable(result):
    if result.empty or result.shape[1] < 2:
        return None
    label = result.columns[0]
    number = next((column for column in result.columns[1:] if pd.api.types.is_numeric_dtype(result[column])), None)
    if number is None:
        return None
    return result.set_index(label)[number]


# Display the SQL query interface for the reviewed postings
def render_query(reviews):
    st.subheader("Query the backlog")
    st.write(
        "The reviewed postings are a table named dataset. Edit the SQL and run it. "
        "A bar chart is drawn when the result has a label column and a number column."
    )
    choice = st.selectbox("Example query", list(QUERIES))
    if st.session_state.get("query_choice") != choice:
        st.session_state.query_choice = choice
        st.session_state.sql_text = QUERIES[choice]
    st.text_area("SQL", height=160, key="sql_text")
    if st.button("Run query"):
        try:
            result = ps.sqldf(st.session_state.sql_text, {"dataset": reviews})
        except Exception as exc:
            st.session_state.sql_error = str(exc)
            st.session_state.sql_result = None
        else:
            st.session_state.sql_error = ""
            st.session_state.sql_result = result
    if st.session_state.get("sql_error"):
        st.error(st.session_state.sql_error)
    result = st.session_state.get("sql_result")
    if isinstance(result, pd.DataFrame):
        st.dataframe(result, hide_index=True)
        chart = chartable(result)
        if chart is not None:
            st.bar_chart(chart)


# Display the upload form, manual inputs, and prediction result
def render_check(model, scores):
    st.subheader("Check a new posting")
    st.write(
        "Enter the listing, or drop a PDF or Word file and the app will fill what it can read. "
        "The box below returns scam = 1 or scam = 0, then the three closest reviewed postings."
    )
    st.caption(f"On held-out reviewed postings this check is right {scores['accuracy']:.1%} of the time.")

    st.subheader("Drop a job posting")
    uploaded = st.file_uploader(
        "Drop a job posting",
        type=["pdf", "docx"],
        help="Accepts a PDF or a .docx file. The verdict uses the fields below, after you check them.",
        label_visibility="collapsed",
    )
    if uploaded is not None:
        data = uploaded.getvalue()
        digest = hashlib.sha256(data).hexdigest()
        if digest != st.session_state.get("upload_digest"):
            try:
                found, notes = extract_upload(uploaded.name, data)
            except Exception as exc:
                notes = [f"Could not read that file ({exc}). Fill the fields below instead."]
                found = {}
            apply_extracted(found, notes)
            st.session_state.upload_digest = digest
    if st.session_state.filled_note:
        st.info(st.session_state.filled_note)

    st.subheader("Or input manually")
    left, right = st.columns(2)
    with left:
        st.number_input("Minimum salary", min_value=0.0, step=100.0, key="salary_min", disabled=st.session_state.no_pay_range)
        st.number_input("Maximum salary", min_value=0.0, step=100.0, key="salary_max", disabled=st.session_state.no_pay_range)
        st.checkbox("No pay range was published", key="no_pay_range")
        st.number_input("Experience required (years)", min_value=0.0, max_value=80.0, step=0.5, key="required_experience_years")
        st.number_input("Skills listed", min_value=0, max_value=100, step=1, key="skills_listed")
        st.number_input("Capital-letter ratio", min_value=0.0, max_value=1.0, step=0.001, format="%.3f", key="allcaps_ratio")
        st.toggle("Asks for a trial fee, deposit, or payment details", key="mentions_trial_or_payment")
    with right:
        st.number_input("Hour posted (0-23)", min_value=0, max_value=23, step=1, key="posting_hour")
        st.number_input("Company age on the board (months)", min_value=0.0, max_value=2000.0, step=1.0, key="company_age_months")
        st.number_input("Company size", min_value=1, max_value=100000, step=1, key="company_size")

    if st.button("Check this posting", type="primary"):
        posting = posting_frame()
        if posting is None:
            st.session_state.verdict = None
            st.session_state.probability = None
            st.session_state.closest_table = None
            st.session_state.pattern_body = ""
            st.session_state.result_text = (
                "Fill in every field. If the posting published no pay range, tick that box."
            )
        else:
            verdict, probability, closest = model.score(posting)
            lines = pattern_lines(model, posting)
            st.session_state.result_text = result_text(verdict, probability, closest, lines)
            st.session_state.verdict = verdict
            st.session_state.probability = probability
            st.session_state.closest_table = neighbor_table(closest, model.columns)
            st.session_state.pattern_body = "\n".join(lines)

    verdict_panel(st.session_state.verdict, st.session_state.probability)

    if st.session_state.verdict is None:
        st.text_area("Result", key="result_text", height=280)
        return

    st.text_area("Result", key="result_text", height=280)
    st.subheader("Three closest reviewed postings")
    st.caption("scam = 1 means the trust and safety team confirmed that reviewed posting was a scam.")
    st.dataframe(st.session_state.closest_table, hide_index=True)
    st.subheader("Feature pattern")
    st.text(st.session_state.pattern_body)


# Set up missing Streamlit session values and clear old form versions
def ensure_state(defaults):
    # Version 2 clears the old median values that used to sit in the form.
    if st.session_state.get("form_version") != 2:
        for key, value in defaults.items():
            st.session_state[key] = value
        st.session_state.form_version = 2
    else:
        for key, value in defaults.items():
            st.session_state.setdefault(key, value)
    st.session_state.setdefault("filled_note", "")
    st.session_state.setdefault("result_text", "Check a posting to see the verdict.")
    st.session_state.setdefault("verdict", None)
    st.session_state.setdefault("probability", None)
    st.session_state.setdefault("closest_table", None)
    st.session_state.setdefault("pattern_body", "")


# Build the Streamlit application pages
def render():
    st.set_page_config(page_title="Job posting check", layout="wide")
    if not train.MODEL_PATH.exists():
        st.error("model.joblib is missing. In this folder, run: python train.py")
        st.stop()
    model = cached_saved_model(train.MODEL_PATH.stat().st_mtime)
    reviews = model.reviews
    scores = model.scores
    ensure_state(form_defaults())

    st.title("Job posting check")
    st.write(
        "A moderation desk for fake job listings. People applying for work run into postings "
        "that collect personal details or ask for money. Check a new listing against the saved "
        "model, explore the reviewed backlog, and query it with SQL."
    )
    check_tab, backlog_tab, query_tab = st.tabs(["Check a posting", "Backlog", "Query"])
    with check_tab:
        render_check(model, scores)
    with backlog_tab:
        render_backlog(reviews, scores, model.columns)
    with query_tab:
        render_query(reviews)


# Check whether the file is being run directly by Streamlit
def running_in_streamlit():
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx

        return get_script_run_ctx() is not None
    except Exception:
        return False


if running_in_streamlit():
    render()
